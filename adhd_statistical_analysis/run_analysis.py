"""Regenerate all analysis outputs from immutable source files and published summaries."""
from __future__ import annotations

import argparse
import hashlib
import importlib.metadata
import json
import os
import platform
from datetime import datetime, timezone
from pathlib import Path

ROOT = Path(__file__).resolve().parent
os.environ.setdefault("MPLCONFIGDIR", str(ROOT / ".mplconfig"))
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import openpyxl
import pandas as pd
import seaborn as sns
import statsmodels.api as sm
from scipy import stats
from statsmodels.nonparametric.smoothers_lowess import lowess
from statsmodels.stats.outliers_influence import variance_inflation_factor

from statistics_helpers import holm, mean_ci, pearson_with_ci, spearman_with_ci, welch_from_summary

OUT = ROOT / "outputs"
IA = [1, 2, 3, 4, 7, 8, 9, 10, 11]
HI = [5, 6, 12, 13, 14, 15, 16, 17, 18]
LABELS = {
    "asrs_total": "ASRS total (0–72)", "inattention": "ASRS inattention (0–36)",
    "hyperactivity": "ASRS hyperactivity/impulsivity (0–36)",
    "aas_complete": "Stored baseline AAS total (9–45), complete items",
    "aas_as_released": "Stored baseline AAS, including partial-item totals",
    "age": "Age (years)", "matric_mark": "Grade 12 score (%)", "nbt_ave": "NBT average (%)",
    "bdi1_total": "BDI-II total (0–63)", "bai1_total": "Stored BAI total (discrepant; descriptive only)",
    "audit1_total": "Released 9-item AUDIT total (descriptive only)",
    "aas_change": "Released AAS change field (unverified; descriptive only)",
}
COLORS = ["#275D8C", "#B75D45"]


def dump_json(path, value):
    path.write_text(json.dumps(value, indent=2, default=lambda obj: obj.item() if isinstance(obj, np.generic) else str(obj)))


def describe(values, source, variable, group, units, total_n, caveat=""):
    x = pd.Series(values).dropna().astype(float)
    n = len(x)
    mean, sd = x.mean(), x.std(ddof=1)
    low, high = mean_ci(n, mean, sd) if n > 1 else (np.nan, np.nan)
    q1, q3 = x.quantile([0.25, 0.75]) if n else (np.nan, np.nan)
    return dict(source=source, variable=variable, label=LABELS.get(variable, variable), group=group,
                units=units, data_type="participant-level released data", n=n, total_n=total_n,
                missing_n=total_n-n, missing_percent=100*(total_n-n)/total_n, mean=mean,
                sample_sd=sd, median=x.median(), q1=q1, q3=q3, iqr=q3-q1,
                minimum=x.min(), maximum=x.max(), skewness=stats.skew(x, bias=False) if n > 2 else np.nan,
                mean_ci_low=low, mean_ci_high=high, unavailable_statistics="", caveat=caveat)


def israel_analysis():
    summaries = pd.read_csv(ROOT / "data" / "israel_table3.csv")
    tests, descriptives = [], []
    for _, row in summaries.iterrows():
        low, high = mean_ci(row.n, row["mean"], row.sd)
        descriptives.append(dict(source="Israel: Sarid & Lipka (2023), Table 3", variable=row.outcome,
            label=row.outcome.replace("_", " "), group=row.group, units=row.units,
            data_type="published rounded summaries only", n=row.n, total_n=np.nan,
            missing_n=np.nan, missing_percent=np.nan, mean=row["mean"], sample_sd=row.sd,
            median=np.nan, q1=np.nan, q3=np.nan, iqr=np.nan, minimum=np.nan, maximum=np.nan, skewness=np.nan,
            mean_ci_low=low, mean_ci_high=high,
            unavailable_statistics="Missing N/%, medians, quartiles, IQR, min/max, skewness, outliers, Shapiro and Levene: not computable from summary statistics",
            caveat="Self-reported diagnosis with author-described verification; self-reported adjustment during COVID-19 remote learning"))
    for outcome, rows in summaries.groupby("outcome", sort=False):
        a = rows.loc[rows.group == "ADHD only"].iloc[0]
        b = rows.loc[rows.group == "No LD/ADHD"].iloc[0]
        result = welch_from_summary(a.n, a["mean"], a.sd, b.n, b["mean"], b.sd)
        reference = stats.ttest_ind_from_stats(a["mean"], a.sd, a.n, b["mean"], b.sd, b.n, equal_var=False)
        if not np.isclose(result["p_raw"], reference.pvalue, rtol=1e-12):
            raise AssertionError("Independent SciPy Welch cross-check failed")
        tests.append(dict(source="Israel: Sarid & Lipka (2023), Table 3", population="Israeli Hebrew-speaking undergraduates, spring 2020 remote learning",
            role="primary", family="Israel_primary_m2", method="Welch independent two-sample t, two-sided",
            predictor="Self-reported ADHD-only diagnosis (author verification)", outcome=outcome,
            group1="ADHD only", group0="No LD/ADHD", n=int(a.n+b.n), units=a.units,
            estimate=result["difference"], effect_size=result["hedges_g"], effect_size_type="Hedges g (ADHD minus comparison)",
            ci_method="Satterthwaite t, marginal 95% mean-difference CI", alternative="two-sided",
            diagnostics="Individual-level diagnostics not computable from summary statistics; Welch allows unequal variances",
            caveat="Rounded published summaries; unadjusted group contrast, no covariate or matching adjustment; self-reported outcomes, not GPA or observed social ability",
            **result))
    adjusted = holm([row["p_raw"] for row in tests])
    for row, value in zip(tests, adjusted):
        row.update(p_holm=value, supported_at_05=bool(value < .05 and row["estimate"] < 0),
                   effect_direction="Lower adjustment in ADHD-only group")
    return summaries, tests, descriptives


def clean_south_africa(raw_path):
    raw = pd.read_excel(raw_path, sheet_name="Sheet1", engine="openpyxl", keep_default_na=True)
    wb = openpyxl.load_workbook(raw_path, data_only=False, read_only=False)
    inventory, dictionary = [], []
    for col in raw.columns:
        x = raw[col]
        numeric = pd.to_numeric(x, errors="coerce")
        inventory.append(dict(column=col, dtype=str(x.dtype), n=len(x), nonmissing_n=int(x.notna().sum()),
            missing_n=int(x.isna().sum()), missing_percent=float(x.isna().mean()*100),
            distinct_nonmissing=int(x.nunique()), numeric_n=int(numeric.notna().sum()),
            numeric_min=numeric.min(), numeric_max=numeric.max()))
        meaning, status, direction = "Source header only; not used in inference", "descriptive inventory only", "not established"
        if col.startswith("asrs1_item_"):
            number = int(col.rsplit("_", 1)[1]); meaning = f"Time 1 ASRS item {number}; {'inattention' if number in IA else 'hyperactivity/impulsivity'}; 0–4"
            status, direction = "verified in Mohamad et al. Methods; no reverse scoring", "higher = more symptoms"
        elif col == "asrs1_total.x":
            meaning, status, direction = "Time 1 ASRS total, 18 items, 0–72", "verified against item sum in all rows", "higher = more symptoms"
        elif col.startswith("aas1_item_"):
            meaning, status, direction = "Baseline AAS item, 1–5", "identity/order and prior reverse coding not independently documented", "item dependent; total intended higher = better"
        elif col == "aas1_total":
            meaning, status, direction = "Stored baseline AAS item sum; nine items imply range 9–45", "provisional numerical audit only; incomplete totals set missing in analysis; reverse coding unresolved", "article: higher = better adjustment; actual coding uncertain"
        elif col == "asrs1_total.y":
            meaning, status = "Ambiguous merge artifact: values largely track BAI, not ASRS", "excluded from inference; not treated as a second ASRS wave"
        elif col == "aas_change":
            meaning, status = "Field named AAS change; may relate to article's mixed-model slopes", "descriptive only; exact construction, wave data and IDs unavailable"
        elif col == "psy1004_grade":
            meaning, status, direction = "Header indicates a single psychology-course grade; differs from paper's weighted GPA", "excluded from GPA inference; no exact GPA mapping available", "higher grade intended better; exact metadata unavailable"
        elif col in ["age", "matric_mark", "nbt_ave"]:
            meaning = {"age":"Age in years at Time 1", "matric_mark":"Grade 12 mark; average home-language and five highest other subjects, excluding Life Orientation, 0–100%", "nbt_ave":"Author-supplied NBT average; article specifies literacy, quantitative literacy and mathematics, 0–100%"}[col]
            status, direction = "mapped using article and matching descriptive statistics; use supplied NBT average", "higher academic marks = better; age not an outcome"
        elif col == "bdi1_total" or col.startswith("bdi1_item_"):
            meaning, status, direction = "Time 1 BDI-II; items 0–3, total 0–63", "total verified against 21-item sum", "higher = more depressive symptoms"
        elif col == "bai1_total" or col.startswith("bai1_item_"):
            meaning, status, direction = "Time 1 BAI; items 0–3, total intended 0–63", "stored total conflicts with sum in 61 rows; excluded as regression covariate", "higher = more anxiety symptoms"
        elif col == "audit1_total" or col.startswith("audit1_item_"):
            meaning, status, direction = "Released modified 9-item AUDIT", "stored sum verified; conventional 10-item scoring not established; descriptive only", "higher = more risky alcohol use, per article"
        elif col.startswith("nbt_"):
            meaning, status = "NBT-related source field; exact field-specific construction not independently established", "not used as a substitute for nbt_ave"
        dictionary.append(dict(column=col, meaning=meaning, verification_status=status, direction=direction,
            source="Rousseau et al. 2024 workbook Sheet1; Mohamad et al. 2025 Methods"))
    pd.DataFrame(inventory).to_csv(OUT / "column_inventory.csv", index=False)
    pd.DataFrame(dictionary).to_csv(OUT / "data_dictionary.csv", index=False)
    validation = {}
    for scale, count, high in [("asrs",18,4), ("aas",9,5), ("bdi",21,3), ("bai",21,3), ("audit",9,4)]:
        cols = [f"{scale}1_item_{i}" for i in range(1, count+1)]
        values = raw[cols].apply(pd.to_numeric, errors="raise")
        bad = values.notna() & ((values < (1 if scale == "aas" else 0)) | (values > high) | (values % 1 != 0))
        if bad.any().any():
            raise ValueError(f"Unexpected {scale} item range: review before analysis")
        total_col = f"{scale}1_total" + (".x" if scale == "asrs" else "")
        complete = values.notna().all(axis=1)
        strict_sum = values.sum(axis=1, min_count=count)
        validation[scale] = dict(complete_item_n=int(complete.sum()), incomplete_item_n=int((~complete).sum()),
            stored_total_n=int(raw[total_col].notna().sum()),
            item_sum_mismatches_complete=int((~np.isclose(strict_sum[complete], raw.loc[complete,total_col])).sum()),
            stored_total_equals_skipna_sum_n=int(np.isclose(values.sum(axis=1),raw[total_col]).sum()))
    if validation["asrs"]["item_sum_mismatches_complete"]:
        raise ValueError("ASRS total no longer matches items; review changed source")
    data = pd.DataFrame({"source_excel_row": np.arange(2, len(raw)+2)})
    data["asrs_total"] = raw[[f"asrs1_item_{i}" for i in range(1,19)]].sum(axis=1, min_count=18)
    data["inattention"] = raw[[f"asrs1_item_{i}" for i in IA]].sum(axis=1, min_count=9)
    data["hyperactivity"] = raw[[f"asrs1_item_{i}" for i in HI]].sum(axis=1, min_count=9)
    complete = raw[[f"aas1_item_{i}" for i in range(1,10)]].notna().all(axis=1)
    data["aas_complete"] = raw.aas1_total.where(complete)
    data["aas_as_released"] = raw.aas1_total
    for column in ["age","matric_mark","nbt_ave","bdi1_total","bai1_total","audit1_total","aas_change"]:
        data[column] = pd.to_numeric(raw[column], errors="raise")
    data.to_csv(OUT / "prepared_analysis_data.csv", index=False)
    formula_errors = [(sheet.title, cell.coordinate, str(cell.value)) for sheet in wb for row in sheet for cell in row
                      if cell.data_type == "e" or (cell.data_type == "f" and "#REF!" in str(cell.value))]
    validation.update(n_rows=len(raw), n_columns=len(raw.columns), exact_duplicate_rows=int(raw.duplicated().sum()),
        sheet_dimensions={s.title:{"rows":s.max_row,"columns":s.max_column} for s in wb},
        participant_id_available=False, wave2_or_wave3_items_available=False, b1_b2_present=False,
        formula_errors=formula_errors,
        asrs_y_matches_bai_stored_n=int((raw["asrs1_total.y"]==raw.bai1_total).sum()),
        asrs_y_matches_bai_item_sum_n=int(np.isclose(raw["asrs1_total.y"],raw[[f"bai1_item_{i}" for i in range(1,22)]].sum(axis=1)).sum()),
        aas_reverse_coding="Unresolved: author total is untransformed sum; original AAS appendix marks items 2 and 3 reversed; workbook gives no question wording or recoding history",
        missing_codes="Numeric analysis fields: Excel blanks only. Text history fields contain na/n/a/none/not applicable; not used or exported.",
        outliers_sheet_policy="10 headerless ancillary rows; no IDs, unresolved provenance and broken external formula references. Not appended to canonical 506-row Sheet1; no new participant exclusion based on observed results.")
    dump_json(OUT / "data_validation.json", validation)
    return raw, data, validation


def summary_plot(summaries, result):
    rows = summaries.loc[summaries.outcome == result["outcome"]].set_index("group").loc[["ADHD only", "No LD/ADHD"]]
    sns.set_theme(style="whitegrid", font_scale=1.05)
    fig, ax = plt.subplots(figsize=(8.8, 5.6))
    for index, (group, row) in enumerate(rows.iterrows()):
        low, high = mean_ci(row.n, row["mean"], row.sd)
        ax.errorbar(index, row["mean"], yerr=[[row["mean"]-low],[high-row["mean"]]], fmt="o", color=COLORS[index], markersize=10, capsize=7, linewidth=2)
        ax.annotate(f"Mean {row['mean']:.2f}\nn={int(row.n)}", (index, high), xytext=(0, 12), textcoords="offset points", ha="center", fontsize=11)
    name = result["outcome"].replace("_", " ").capitalize()
    ax.set(title=f"{name} during COVID-19 remote learning", ylabel="Adapted SACQ score (1–9; higher = better)",
        xticks=[0,1], xticklabels=["ADHD only", "No LD/ADHD"], xlim=(-.55,1.55), ylim=(1,9))
    ax.set_yticks(range(1,10))
    ax.grid(axis="x", visible=False)
    fig.text(.5, .105, f"ADHD minus comparison: {result['difference']:.2f} points (95% CI {result['ci_low']:.2f} to {result['ci_high']:.2f})\n"
        f"Welch p={result['p_raw']:.5f}; Holm p={result['p_holm']:.5f}; Hedges' g={result['hedges_g']:.2f}", ha="center", fontsize=10)
    fig.text(.5, .025, "Sarid & Lipka (2023), Table 3. Published rounded summaries; points and marginal 95% mean CIs.\n"
        "Self-reported diagnoses with author verification; self-reported adjustment. No individual data.", ha="center", fontsize=9, color="#444444")
    fig.subplots_adjust(bottom=.27, top=.89, left=.12, right=.97)
    filename = "academic_results.png" if result["outcome"] == "academic_adjustment" else "social_results.png"
    fig.savefig(OUT / filename, dpi=180)
    plt.close(fig)


def raw_plots(data):
    columns = ["asrs_total", "inattention", "hyperactivity", "aas_complete"]
    sns.set_theme(style="whitegrid", font_scale=.9)
    fig, axes = plt.subplots(4, 3, figsize=(13, 13))
    diagnostics = []
    for index, col in enumerate(columns):
        x = data[col].dropna()
        sns.histplot(x=x, bins=20, color=COLORS[0], ax=axes[index,0])
        axes[index,0].set(xlabel=LABELS[col], ylabel="Count", title=f"n={len(x)}")
        sns.boxplot(x=x, color="#AFC9DE", ax=axes[index,1], whis=1.5)
        axes[index,1].set(xlabel=LABELS[col], title="Tukey boxplot; all observations retained")
        stats.probplot(x, dist="norm", plot=axes[index,2])
        axes[index,2].set_title("Normal Q–Q plot")
        sw = stats.shapiro(x)
        q1, q3 = x.quantile([.25,.75]); iqr=q3-q1
        diagnostics.append(dict(variable=col, n=len(x), shapiro_w=sw.statistic, p_raw=sw.pvalue,
            p_adjusted="not applicable: diagnostic only", tukey_outlier_n=int(((x<q1-1.5*iqr)|(x>q3+1.5*iqr)).sum()),
            exclusion_rule="none", interpretation="Large N and bounded scores; Shapiro does not determine test selection"))
    fig.suptitle("South African baseline distributions\nAAS is the stored score; reverse coding remains unresolved", fontsize=15)
    fig.tight_layout(rect=(0,0,1,.95))
    fig.savefig(OUT / "south_africa_distributions.png", dpi=150)
    plt.close(fig)
    pd.DataFrame(diagnostics).to_csv(OUT / "distribution_diagnostics.csv", index=False)
    pair=data.dropna(subset=["aas_complete"])
    fig, axes=plt.subplots(1,3,figsize=(14,4.8),sharey=True)
    for ax,col in zip(axes,columns[:3]):
        ax.scatter(pair[col],pair.aas_complete,s=15,alpha=.28,color=COLORS[0],edgecolor="none")
        smooth=lowess(pair.aas_complete,pair[col],frac=.65)
        ax.plot(smooth[:,0],smooth[:,1],color=COLORS[1],linewidth=2,label="LOWESS descriptive trend")
        ax.set(xlabel=LABELS[col],ylabel="Stored baseline AAS total",title=f"Complete AAS items: n={len(pair)}")
    axes[0].legend(fontsize=8)
    fig.suptitle("Symptoms and the released AAS score: provisional score audit, not GPA",fontsize=14)
    fig.text(.5,.015,"No recoding, jitter, imputation or outlier deletion. Scale coding needs author clarification.",ha="center",fontsize=9)
    fig.tight_layout(rect=(0,.06,1,.92))
    fig.savefig(OUT / "south_africa_associations.png",dpi=170)
    plt.close(fig)


def south_africa_analysis(raw, data):
    tests=[]
    for method in ["Pearson", "Spearman"]:
        family=[]
        for predictor in ["asrs_total", "inattention", "hyperactivity"]:
            pair=data[[predictor,"aas_complete"]].dropna()
            calc=(pearson_with_ci if method=="Pearson" else spearman_with_ci)(pair[predictor],pair.aas_complete)
            result=dict(source="South Africa: Rousseau et al. (2024), Sheet1",population="Convenience sample of first-time undergraduates in introductory psychology, Cape Town, 2023",
                role="exploratory provisional score audit",family=f"SA_provisional_{method.lower()}_m3",method=f"{method} correlation, two-sided",
                predictor=predictor,outcome="aas_complete",group1="continuous ASRS symptoms",group0="not applicable: no diagnosis/control grouping",
                n=len(pair),n1=np.nan,n0=np.nan,units="ASRS sum points versus stored baseline AAS points",alternative="two-sided",
                effect_size=calc["estimate"],effect_size_type="Pearson r" if method=="Pearson" else "Spearman rho",
                diagnostics="See distributions and scatterplots; no data-driven exclusions. Bounded/discrete scores; rank sensitivity provided.",
                caveat="Stored AAS reverse coding unresolved; this is numerical reproduction only, not verified academic-adjustment inference; released ASRS may contain author imputations",**calc)
            family.append(result)
        for result,pvalue in zip(family,holm([t["p_raw"] for t in family])):
            result.update(p_holm=pvalue,supported_at_05="not adjudicated: unresolved AAS coding",
                effect_direction="Higher symptoms accompany lower stored AAS score" if result["estimate"]<0 else "Higher symptoms accompany higher stored AAS score")
        tests.extend(family)
    # Historical numerical reconstruction only. This rejected construction is never an inferential alternative.
    wrong=raw[[f"asrs1_item_{i}" for i in range(1,10)]].sum(axis=1,min_count=9)
    complete=data.aas_complete.notna()
    audit_by_predictor={result["predictor"]:result for result in tests if result["family"]=="SA_provisional_pearson_m3"}
    historical=dict(documented_inattention_complete_n=int(complete.sum()),
        documented_inattention_r=float(np.corrcoef(data.loc[complete,"inattention"],data.loc[complete,"aas_complete"])[0,1]),
        documented_inattention_p_raw=audit_by_predictor["inattention"]["p_raw"],
        documented_inattention_p_holm_m3=float(audit_by_predictor["inattention"]["p_holm"]),
        documented_total_asrs_r=float(np.corrcoef(data.loc[complete,"asrs_total"],data.loc[complete,"aas_complete"])[0,1]),
        documented_total_asrs_n=int(complete.sum()),
        documented_total_asrs_p_raw=audit_by_predictor["asrs_total"]["p_raw"],
        documented_total_asrs_p_holm_m3=float(audit_by_predictor["asrs_total"]["p_holm"]),
        second_coefficient_assessment="The correctly scored ASRS total association rounds to -0.14; this is a different exposure from inattention. The previous Holm p=.015 cannot be verified without the earlier family and code.",
        rejected_first_nine_sum_r=float(np.corrcoef(wrong[complete],data.loc[complete,"aas_complete"])[0,1]),
        rejected_reason="First nine ASRS items mix inattention and hyperactivity/impulsivity. Source specifies 1–4 and 7–11.",
        no_p_value_for_rejected_score=True,
        previous_claims=[{"r":-.205,"n":488,"p_raw":5.12e-6},{"r":-.14,"p_holm":.015}],
        limitation="No prior analysis script, full sample definition, adjustment family or presentation available; provenance of the second claim cannot be established.")
    dump_json(OUT/"historical_reconciliation.json",historical)
    # Prespecified, conditional exploratory regression audits, not GPA models.
    models=[];coefficients=[];vifs=[];model_info=[]
    for name,predictors in [("AAS_base",["asrs_total","age","matric_mark","nbt_ave"]),
                             ("AAS_plus_BDI",["asrs_total","age","matric_mark","nbt_ave","bdi1_total"])]:
        sample=data[["source_excel_row","aas_complete",*predictors]].dropna()
        X=sm.add_constant(sample[predictors],has_constant="add")
        ordinary=sm.OLS(sample.aas_complete,X).fit()
        model=ordinary.get_robustcov_results(cov_type="HC3",use_t=True)
        ci=model.conf_int(); pos=list(X.columns).index("asrs_total")
        for i,col in enumerate(X.columns):
            coefficients.append(dict(model=name,term=col,n=len(sample),estimate=model.params[i],se_hc3=model.bse[i],ci_low=ci[i,0],ci_high=ci[i,1],
                df=model.df_resid,r_squared=model.rsquared,adjusted_r_squared=model.rsquared_adj,
                p_raw=model.pvalues[i] if col=="asrs_total" else np.nan,
                p_holm=np.nan,role="ASRS coefficient prespecified" if col=="asrs_total" else "nuisance parameter estimate only; not tested",
                caveat="Provisional stored AAS score; scoring unresolved; marginal HC3 t interval"))
            if col!="const":vifs.append(dict(model=name,term=col,vif=variance_inflation_factor(X.to_numpy(),i)))
        influence=ordinary.get_influence();cook=influence.cooks_distance[0]
        info=dict(model=name,n=len(sample),r_squared=model.rsquared,adjusted_r_squared=model.rsquared_adj,
            df_resid=model.df_resid,covariates=predictors,max_cooks_distance=float(cook.max()),cooks_above_4_over_n=int((cook>4/len(sample)).sum()),
            max_leverage=float(influence.hat_matrix_diag.max()),max_abs_studentized_residual=float(np.max(np.abs(influence.resid_studentized_internal))),
            covariance="HC3 robust, t reference",records_deleted_for_influence=0)
        model_info.append(info)
        models.append(dict(source="South Africa: Rousseau et al. (2024), Sheet1",population="Cape Town first-year student baseline complete AAS cases",
            role="exploratory provisional score audit",family="SA_provisional_regression_ASRS_m2",method="OLS with HC3 robust SE and two-sided t reference",
            predictor="asrs_total",outcome="aas_complete",model=name,n=len(sample),n1=np.nan,n0=np.nan,
            group1="continuous ASRS symptoms",group0="not applicable",units="Stored AAS points per 1 ASRS point",
            estimate=model.params[pos],statistic=model.tvalues[pos],df=model.df_resid,p_raw=model.pvalues[pos],ci_low=ci[pos,0],ci_high=ci[pos,1],
            ci_method="HC3 robust t, marginal 95%",effect_size=model.params[pos],effect_size_type="unstandardized coefficient",r_squared=model.rsquared,
            alternative="two-sided",covariates=", ".join(predictors[1:]),diagnostics="See regression_vif.csv, regression_diagnostics.json and residual plots; no influential-case deletions",
            caveat="Not a GPA regression. Concurrent AAS and ASRS; reverse scoring unresolved; BDI addition is explanatory sensitivity, not causal adjustment"))
        if name=="AAS_base":
            fig,axes=plt.subplots(1,3,figsize=(13,4.4))
            axes[0].scatter(ordinary.fittedvalues,ordinary.resid,alpha=.35,s=15);axes[0].axhline(0,color=COLORS[1]);axes[0].set(xlabel="Fitted stored AAS",ylabel="Residual",title="Residuals versus fitted")
            stats.probplot(ordinary.resid,dist="norm",plot=axes[1]);axes[1].set_title("Residual normal Q–Q")
            axes[2].scatter(sample.source_excel_row,cook,alpha=.5,s=15);axes[2].axhline(4/len(sample),color=COLORS[1],linestyle="--");axes[2].set(xlabel="Source Excel row (not participant ID)",ylabel="Cook's distance",title="Influence; all rows retained")
            fig.suptitle("Provisional AAS regression diagnostics; HC3 inference")
            fig.tight_layout();fig.savefig(OUT/"regression_diagnostics.png",dpi=150);plt.close(fig)
    for result,value in zip(models,holm([r["p_raw"] for r in models])):
        result.update(p_holm=value,supported_at_05="not adjudicated: unresolved AAS coding",effect_direction="Negative ASRS coefficient" if result["estimate"]<0 else "Positive ASRS coefficient")
        for coeff in coefficients:
            if coeff["model"]==result["model"] and coeff["term"]=="asrs_total":coeff["p_holm"]=value
    pd.DataFrame(coefficients).to_csv(OUT/"regression_coefficients.csv",index=False)
    pd.DataFrame(vifs).to_csv(OUT/"regression_vif.csv",index=False)
    dump_json(OUT/"regression_diagnostics.json",model_info)
    tests.extend(models)
    return tests,historical


def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--raw-file",type=Path,default=ROOT/"data/raw/Student Mental Health Article Data.xlsx")
    args=parser.parse_args();OUT.mkdir(exist_ok=True,parents=True)
    sns.set_theme(style="whitegrid")
    summaries,tests,descriptives=israel_analysis()
    for result in tests:summary_plot(summaries,result)
    validation={"status":"pending raw data"};historical={};data=None
    raw_path=args.raw_file
    if not raw_path.exists():
        candidate=ROOT.parent/"Student Mental Health Article Data.xlsx"
        if candidate.exists():raw_path=candidate
    if raw_path.exists():
        raw,data,validation=clean_south_africa(raw_path)
        more_tests,historical=south_africa_analysis(raw,data)
        tests.extend(more_tests)
        for col in LABELS:
            caveat="Stored AAS coding unresolved; numerical score description only" if col.startswith("aas") else "Released data may include author-side preprocessing or imputations; no imputation added here"
            if col in ["bai1_total","audit1_total"]:caveat+="; source scoring limitations, not used in inference"
            descriptives.append(describe(data[col],"South Africa: Rousseau et al. (2024), Sheet1",col,"All released baseline participants",LABELS[col],len(data),caveat))
        raw_plots(data)
    all_tests=pd.DataFrame(tests)
    all_tests.to_csv(OUT/"statistical_tests.csv",index=False,float_format="%.12g")
    pd.DataFrame(descriptives).to_csv(OUT/"descriptive_statistics.csv",index=False,float_format="%.12g")
    run_info=dict(seed=42,bootstrap_resamples=5000,alpha=.05,alternative="two-sided",analysis_date="2026-10-20",
        run_timestamp_utc=datetime.now(timezone.utc).isoformat(),python=platform.python_version(),
        packages={name:importlib.metadata.version(name) for name in ["pandas","numpy","scipy","statsmodels","matplotlib","seaborn","openpyxl","beautifulsoup4","pytest"]},
        source_hashes={str(path.relative_to(ROOT.parent)):hashlib.sha256(path.read_bytes()).hexdigest() for path in [ROOT.parent/"Student Mental Health Article Data.xlsx",ROOT.parent/"intro_analysis.ipynb",ROOT.parent/"ADHD_Codex_Statistical_Analysis_Prompt.md",raw_path] if path.exists()},
        raw_available=data is not None)
    dump_json(OUT/"run_metadata.json",run_info)
    from write_report import write_documents
    write_documents(summaries,all_tests,pd.DataFrame(descriptives),validation,historical,run_info)
    print(all_tests.loc[all_tests.role=="primary",["outcome","estimate","ci_low","ci_high","p_raw","p_holm","effect_size"]].to_string(index=False))
    print(f"Generated {len(list(OUT.iterdir()))} outputs in {OUT}")


if __name__=="__main__":main()
