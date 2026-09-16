"""Build the six-slide, editable ADHD topic-proposal deck and its original chart.

Run from any directory using the analysis project's Python environment.
Only source-verified, published group means are plotted. No support exposure is inferred.
"""
from pathlib import Path
import csv
import json
import os
import re
import argparse

import plotly.graph_objects as go
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION, XL_TICK_MARK
from pptx.chart.data import CategoryChartData
from pptx.oxml.xmlchemy import OxmlElement

ROOT = Path(__file__).resolve().parent
ASSETS = ROOT / 'assets'
CHARTS = ROOT / 'charts'
ACCENT = '6058E8'
INK = '18212F'
MUTED = '596373'
LIGHT = 'F3F2FE'
GRAY = 'B7BECA'
LINE = 'DCDDE4'
FONT = 'Inter'
WIDTH, HEIGHT = 13.333333, 7.5

GPA = [
    {'semester':'Fall', 'ADHD':2.91, 'Comparison':3.26, 'reported_d':-.48},
    {'semester':'Spring', 'ADHD':2.79, 'Comparison':3.13, 'reported_d':-.41},
]

SOURCES = {
    'alvarez': {'short':'Álvarez-Godos et al., 2023', 'doi':'10.3389/fpsyg.2023.1216692',
        'authors':'Álvarez-Godos, M., Ferreira, C., & Vieira, M.-J.', 'year':'2023',
        'title':'A systematic review of actions aimed at university students with ADHD.',
        'journal':'Frontiers in Psychology', 'volume':'14', 'issue':'', 'pages':'Article 1216692.'},
    'anastopoulos': {'short':'Anastopoulos et al., 2021', 'doi':'10.1037/ccp0000553',
        'authors':'Anastopoulos, A. D., Langberg, J. M., Eddy, L. D., Silvia, P. J., & Labban, J. D.',
        'year':'2021', 'title':'A randomized controlled trial examining CBT for college students with ADHD.',
        'journal':'Journal of Consulting and Clinical Psychology','volume':'89','issue':'1','pages':'21–33.'},
    'blasey': {'short':'Blasey et al., 2023','doi':'10.1177/00332941221078011',
        'authors':'Blasey, J., Wang, C., & Blasey, R.','year':'2023',
        'title':'Accommodation use and academic outcomes for college students with disabilities.',
        'journal':'Psychological Reports','volume':'126','issue':'4','pages':'1891–1909.'},
    'dupaul': {'short':'DuPaul et al., 2017','doi':'10.1111/ldrp.12143',
        'authors':'DuPaul, G. J., Dahlstrom-Hakki, I., Gormley, M. J., Fu, Q., Pinho, T. D., & Banerjee, M.',
        'year':'2017','title':'College students with ADHD and LD: Effects of support services on academic performance.',
        'journal':'Learning Disabilities Research & Practice','volume':'32','issue':'4','pages':'246–256.'},
    'gormley': {'short':'Gormley et al., 2019','doi':'10.1177/1087054715623046',
        'authors':'Gormley, M. J., DuPaul, G. J., Weyandt, L. L., & Anastopoulos, A. D.',
        'year':'2019','title':'First-year GPA and academic service use among college students with and without ADHD.',
        'journal':'Journal of Attention Disorders','volume':'23','issue':'14','pages':'1766–1779.'},
    'nu': {'short':'Nazarbayev University, n.d.', 'authors':'Nazarbayev University.', 'year':'n.d.',
        'title':'Disability & learning needs.', 'url':'https://nu.edu.kz/students/disability-and-special-learning-needs-support/'},
    'romhild': {'short':'Römhild & Hollederer, 2024','doi':'10.1080/08856257.2023.2195074',
        'authors':'Römhild, A., & Hollederer, A.','year':'2024',
        'title':'Effects of disability-related services, accommodations, and integration on academic success of students with disabilities in higher education: A scoping review.',
        'journal':'European Journal of Special Needs Education','volume':'39','issue':'1','pages':'143–166.'},
    'ross': {'short':'Ross et al., 2026','doi':'10.1038/s44184-026-00196-4',
        'authors':'Ross, F., Dommett, E. J., & Byrom, N.','year':'2026',
        'title':'A systematic review of higher education-based interventions to support the mental health and wellbeing of neurodivergent students.',
        'journal':'npj Mental Health Research','volume':'5','issue':'','pages':'Article 14.'},
}

SCRIPTS = [
    """Our proposal examines ADHD-specific academic support for university students in Kazakhstan. The question is whether available support matches students’ everyday learning needs. We will establish why this matters, then consider two possible responses: executive-function coaching and an integrated support system. We are not choosing between them today.""",
    """University study asks students to manage their own learning: plan ahead, sustain attention, organize materials, monitor progress, and coordinate several deadlines. For a student with ADHD, those demands can interact with difficulties in attention and executive functioning. A student may understand the material yet struggle to turn that understanding into timely, organized work.

The reviews by Álvarez-Godos and colleagues and Ross and colleagues support considering tailored approaches, while also showing limitations in the intervention evidence. For Kazakhstan, our working proposition is that specialized, evidence-based ADHD support appears limited and under-researched. We are not claiming that universities have no support at all. Nazarbayev University, for example, documents disability and special learning-needs services.

The unresolved local question is how consistently students can access support designed around ADHD, and whether it is evaluated. This is a proposition to investigate, not a completed national audit. Why should university decision makers take that question seriously?""",
    """This chart uses published means from Gormley and colleagues’ study of first-year students at universities in the United States. The purple bars represent students with ADHD; the gray bars represent the comparison group. In the fall, average GPA was 2.91 compared with 3.26. In the spring, it was 2.79 compared with 3.13. The gap was therefore about one-third of a GPA point in each semester.

The chart starts at zero to keep the difference in proportion. These are group averages: they do not mean that every student with ADHD performs worse. The comparison is observational, so it cannot establish that ADHD alone caused the gap, and it does not estimate the situation in Kazakhstan.

It also does not show that receiving any generic service improves GPA. The study did not find an independent relationship between typical service use and GPA. The implication is to investigate which forms of support address specific needs.""",
    """The first possible solution is ADHD-specific executive-function coaching. A proposed weekly routine would help students plan priorities, break assignments into manageable tasks, estimate time, agree on concrete actions, and review what happened. The goal is to help students use strategies consistently in their own courses, rather than simply give them more subject teaching.

There are two relevant kinds of evidence. DuPaul and colleagues followed service use at a specialist campus over five years; coaching use was associated with academic gains. Because students selected services, this is not a randomized test of coaching.

The ACCESS randomized trial involved 250 students and found improvements in executive functioning and ADHD symptoms. However, ACCESS combined group cognitive-behavioral therapy with individual mentoring, so its results cannot be attributed to coaching alone or treated as a proven GPA gain.

The strength is the direct match to planning and self-regulation needs. The limitation is recurring staff time and training. Our second option considers the infrastructure around this support.""",
    """The second possible solution is an integrated ADHD support system. A student would have a clear access point, appropriate accommodations, connections to coaching or mentoring, and continuing follow-up. Faculty awareness and referral to mental-health care, when needed, would connect these steps. This changes how services work together; coaching is only one component.

Blasey and colleagues studied 1,980 students registered with disability services. Earlier registration and sustained accommodation use were associated with better academic outcomes. This was an observational study of a broad disability population, not an ADHD-only trial. The scoping review by Römhild and Hollederer also found varied results across services and accommodations, so merely creating a service does not guarantee academic improvement.

The strength is continuity from access to ongoing support. The limitation is coordination across services, faculty and advisers. Neither option is the winner yet. The next stage of the research is to compare these approaches in terms of effectiveness, feasibility, cost, and suitability for universities in Kazakhstan.""",
]


def rgb(h):
    return RGBColor.from_string(h)


def box(slide, x, y, w, h, fill='FFFFFF', line=LINE, radius=True):
    s=slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
                            Inches(x), Inches(y), Inches(w), Inches(h))
    if radius:
        s.adjustments[0]=.12
    s.fill.solid();s.fill.fore_color.rgb=rgb(fill)
    if line:
        s.line.color.rgb=rgb(line);s.line.width=Pt(.9)
    else:s.line.fill.background()
    return s


def text(slide, x, y, w, h, content, size=20, bold=False, color=INK, align=PP_ALIGN.LEFT):
    s=slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf=s.text_frame;tf.clear();tf.word_wrap=True
    tf.margin_left=tf.margin_right=0;tf.margin_top=tf.margin_bottom=0
    tf.vertical_anchor=MSO_ANCHOR.TOP
    for i,line in enumerate(content.split('\n')):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph()
        p.text=line;p.alignment=align;p.space_after=Pt(0);p.space_before=Pt(0);p.line_spacing=1.08
        p.font.name=FONT;p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=rgb(color)
    return s


def arrow(slide,x1,y1,x2,y2,color=ACCENT):
    s=slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    s.line.color.rgb=rgb(color);s.line.width=Pt(1.15)
    ln=s._element.spPr.find('{http://schemas.openxmlformats.org/drawingml/2006/main}ln')
    e=OxmlElement('a:tailEnd');e.set('type','triangle');e.set('w','sm');e.set('len','sm');ln.append(e)
    return s


def footer(slide, number, keys, prefix=''):
    text(slide,.65,6.98,11.65,.25,prefix+'('+ '; '.join(SOURCES[k]['short'] for k in keys)+')',10.5,color=MUTED)
    text(slide,12.24,6.97,.43,.3,f'{number:02d}',11,color=MUTED,align=PP_ALIGN.RIGHT)


def header(slide,kicker,title):
    text(slide,.65,.38,12,.25,kicker.upper(),11,bold=True,color=ACCENT)
    text(slide,.65,.83,12,1.12,title,32,bold=True)


def add_notes(slide,script,timing,evidence=''):
    slide.notes_slide.notes_text_frame.text=(f'TARGET: {timing}\n\nSPOKEN SCRIPT\n{script.strip()}'
        +(f'\n\nEVIDENCE / PRESENTER CHECK (not part of spoken timing)\n{evidence}' if evidence else ''))


def chart_assets():
    CHARTS.mkdir(parents=True,exist_ok=True)
    with (CHARTS/'gormley_2019_table2.csv').open('w',newline='') as f:
        w=csv.DictWriter(f,fieldnames=list(GPA[0]));w.writeheader();w.writerows(GPA)
    fig=go.Figure()
    for name,color in [('ADHD',ACCENT),('Comparison',GRAY)]:
        fig.add_bar(name=name,x=[r['semester'] for r in GPA],y=[r[name] for r in GPA],
                    marker_color='#'+color,text=[f'{r[name]:.2f}' for r in GPA],textposition='outside',
                    textfont=dict(size=30,color='#'+INK),cliponaxis=False)
    fig.update_layout(template='simple_white',width=1400,height=860,barmode='group',bargap=.36,bargroupgap=.06,
        font=dict(family=FONT,size=26,color='#'+INK),paper_bgcolor='white',plot_bgcolor='white',
        title=dict(text='<b>First-Year GPA: Students With and Without ADHD</b>',font_size=32,x=.035,y=.97),
        legend=dict(orientation='h',x=.12,y=1.03,font_size=26),
        margin=dict(l=125,r=45,t=145,b=120))
    fig.update_yaxes(title='Mean GPA',range=[0,4],dtick=1,showgrid=True,gridcolor='#E8EAF0',zeroline=False)
    fig.update_xaxes(title=None,showline=False,tickfont_size=29)
    fig.add_annotation(x=0,y=-.17,xref='paper',yref='paper',xanchor='left',showarrow=False,
        text='Source: Data from Gormley et al. (2019), Table 2. Observational U.S. study.',font_size=19)
    fig.write_html(CHARTS/'first_year_gpa.html',include_plotlyjs=True)
    browser=ROOT.parent/'.browser/chrome-mac-arm64/Google Chrome for Testing.app/Contents/MacOS/Google Chrome for Testing'
    if browser.exists():os.environ.setdefault('BROWSER_PATH',str(browser))
    fig.write_image(CHARTS/'first_year_gpa.png',scale=2)
    fig.write_image(CHARTS/'first_year_gpa.svg')


def add_native_gpa_chart(slide):
    data=CategoryChartData();data.categories=[r['semester'] for r in GPA]
    for series in ['ADHD','Comparison']:data.add_series(series,[r[series] for r in GPA])
    chart=slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED,Inches(.67),Inches(2.03),
                                 Inches(8.04),Inches(4.34),data).chart
    chart.has_title=False;chart.has_legend=True
    chart.legend.position=XL_LEGEND_POSITION.TOP;chart.legend.include_in_layout=False
    chart.legend.font.name=FONT;chart.legend.font.size=Pt(18)
    chart.font.name=FONT;chart.font.size=Pt(18)
    plot=chart.plots[0];plot.gap_width=80;plot.overlap=0
    plot.has_data_labels=True
    dl=plot.data_labels;dl.position=XL_LABEL_POSITION.OUTSIDE_END;dl.number_format='0.00'
    dl.font.name=FONT;dl.font.size=Pt(20);dl.font.bold=True;dl.font.color.rgb=rgb(INK)
    for series,color in zip(chart.series,[ACCENT,GRAY]):
        series.format.fill.solid();series.format.fill.fore_color.rgb=rgb(color)
        series.format.line.fill.background()
    y=chart.value_axis;y.minimum_scale=0;y.maximum_scale=4;y.major_unit=1
    y.has_major_gridlines=True;y.major_gridlines.format.line.color.rgb=rgb('E6E8EE')
    y.major_gridlines.format.line.width=Pt(.6)
    y.tick_labels.font.name=FONT;y.tick_labels.font.size=Pt(16);y.tick_labels.font.color.rgb=rgb(MUTED)
    y.format.line.fill.background();y.major_tick_mark=XL_TICK_MARK.NONE;y.minor_tick_mark=XL_TICK_MARK.NONE
    y.has_title=True;y.axis_title.text_frame.text='Mean GPA'
    for p in y.axis_title.text_frame.paragraphs:p.font.name=FONT;p.font.size=Pt(16)
    x=chart.category_axis;x.major_tick_mark=XL_TICK_MARK.NONE;x.minor_tick_mark=XL_TICK_MARK.NONE
    x.tick_labels.font.name=FONT;x.tick_labels.font.size=Pt(19);x.format.line.fill.background()
    return chart


def reference(slide,key,x,y,w,h):
    s=slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h))
    tf=s.text_frame;tf.clear();tf.word_wrap=True
    tf.margin_left=tf.margin_right=0;tf.margin_top=tf.margin_bottom=0
    p=tf.paragraphs[0];p.line_spacing=1.12;p.space_after=Pt(0)
    # Native hanging indent, 0.18 inch, applied to the reference paragraph.
    pr=p._p.get_or_add_pPr();pr.set('marL',str(Inches(.18)));pr.set('indent',str(-Inches(.18)))
    src=SOURCES[key]
    runs=[(f'{src["authors"]} ({src["year"]}). ',False),
          (src['title']+' ',key=='nu')]
    if 'journal' in src:
        runs.extend([(src['journal']+', '+src['volume'],True),
                     ((f'({src["issue"]})' if src['issue'] else '')+', '+src['pages'],False)])
    for content,italic in runs:
        r=p.add_run();r.text=content;r.font.name=FONT;r.font.size=Pt(12)
        r.font.italic=italic;r.font.color.rgb=rgb(INK)
    p2=tf.add_paragraph();p2.line_spacing=1.0;p2.space_before=Pt(3)
    p2._p.get_or_add_pPr().set('marL',str(Inches(.18)))
    r=p2.add_run();url='https://doi.org/'+src['doi'] if 'doi' in src else src['url']
    r.text=url;r.font.name=FONT;r.font.size=Pt(10.7);r.font.color.rgb=rgb(ACCENT);r.hyperlink.address=url


def build(names='[Presenter name(s)]',course='[Course]',date='[Presentation date]'):
    ROOT.mkdir(exist_ok=True);ASSETS.mkdir(exist_ok=True)
    prs=Presentation();prs.slide_width=Inches(WIDTH);prs.slide_height=Inches(HEIGHT)
    prs.core_properties.title='Supporting University Students with ADHD in Kazakhstan'
    prs.core_properties.subject='Topic proposal: two possible academic-support approaches'
    prs.core_properties.author=names if not names.startswith('[') else ''
    prs.core_properties.keywords='ADHD, Kazakhstan, academic support, topic proposal'
    blank=prs.slide_layouts[6]

    # Cover: editable typography and a restrained outline motif.
    s=prs.slides.add_slide(blank)
    text(s,.72,.55,11,.3,'TOPIC PROPOSAL  /  HIGHER EDUCATION',11,bold=True,color=ACCENT)
    text(s,.72,1.60,11.65,2.55,'Supporting University\nStudents with ADHD\nin Kazakhstan',42,bold=True)
    text(s,.76,4.52,10.6,.88,'Addressing the gap in\nADHD-specific academic support',22,color=MUTED)
    box(s,.76,5.74,2.6,.055,fill=ACCENT,line=None,radius=False)
    text(s,.76,6.48,11.2,.45,f'{names}  ·  {course}  ·  {date}',14,color=MUTED)
    add_notes(s,SCRIPTS[0],'~25 seconds')

    # Content 1: mechanism and bounded Kazakhstan framing.
    s=prs.slides.add_slide(blank)
    header(s,'01 / The problem','University demands can amplify\nADHD-related difficulties')
    cards=[(.65,'ADHD-related\ndifficulties','Attention • planning\nOrganization • self-monitoring'),
           (4.87,'University\ndemands','Independent learning\nMultiple deadlines'),
           (9.09,'ADHD-specific\nsupport gap','Availability • consistency\nEvaluation')]
    for i,(x,title,body) in enumerate(cards):
        box(s,x,2.30,3.58,2.24,fill=LIGHT if i==2 else 'FFFFFF',line=ACCENT)
        text(s,x+.23,2.58,3.14,.86,title,22,bold=True,color=ACCENT if i==2 else INK)
        text(s,x+.23,3.64,3.14,.65,body,18,color=MUTED)
    arrow(s,4.30,3.42,4.76,3.42);arrow(s,8.52,3.42,8.98,3.42)
    text(s,9.29,4.68,3.25,.28,'WORKING PROPOSITION TO INVESTIGATE',10,bold=True,color=ACCENT)
    text(s,.72,5.31,11.7,.42,'General support exists—for example, at Nazarbayev University.',20)
    text(s,.72,5.96,11.7,.58,'The research gap: how well does support meet ADHD-specific learning needs?',20,bold=True)
    footer(s,1,['alvarez','ross','nu'])
    add_notes(s,SCRIPTS[1],'~70 seconds',
        'The local gap is a working proposition, not an established national prevalence or service audit. '
        'NU provides a documented counterexample to any claim of no support. International reviews do not establish Kazakhstan-wide availability. '
        'Sources: https://doi.org/10.3389/fpsyg.2023.1216692 ; https://doi.org/10.1038/s44184-026-00196-4 ; '+SOURCES['nu']['url'])

    # Content 2: original, native-editable graph with an embedded Excel data sheet.
    s=prs.slides.add_slide(blank)
    header(s,'02 / Why the problem matters','First-Year GPA: Students\nWith and Without ADHD')
    add_native_gpa_chart(s)
    box(s,9.10,2.58,3.55,2.40,line=ACCENT)
    text(s,9.36,2.94,3.04,.64,'0.34–0.35',34,bold=True,color=ACCENT)
    text(s,9.36,3.82,3.04,.92,'GPA points lower\nin the ADHD group',21)
    text(s,9.15,5.38,3.51,.87,'A consistent gap\nacross both semesters.',20,bold=True)
    text(s,.89,6.53,11.6,.31,'U.S. first-year sample · Observational comparison · Group means do not describe every student',12,color=MUTED)
    text(s,.65,6.98,11.6,.25,'Source: Data from Gormley et al. (2019), Table 2.',10.5,color=MUTED)
    text(s,12.24,6.97,.43,.3,'02',11,color=MUTED,align=PP_ALIGN.RIGHT)
    add_notes(s,SCRIPTS[2],'~75 seconds',
        'Table 2: Fall ADHD M=2.91, SD=.77; comparison M=3.26, SD=.69; reported d=-.48. '
        'Spring ADHD M=2.79, SD=.84; comparison M=3.13, SD=.82; reported d=-.41. '
        'Bars are published descriptive means, not adjusted predictions or a support comparison. '
        'No error bars are invented. Differences .35 and .34 are simple subtraction of those means. '
        'https://doi.org/10.1177/1087054715623046 ; https://pmc.ncbi.nlm.nih.gov/articles/PMC6209537/')

    # Content 3: specific intervention, with evidence distinguished by design.
    s=prs.slides.add_slide(blank)
    header(s,'03 / Possible solution 1','ADHD-specific\nexecutive-function coaching')
    labels=['Weekly\nsupport','Planning','Task\nbreakdown','Time\nmanagement','Accountability','Review']
    for i,label in enumerate(labels):
        x=.65+i*2.07
        box(s,x,2.20,1.69,1.13,fill=LIGHT if i==0 else 'FFFFFF',line=ACCENT)
        text(s,x+.10,2.50,1.49,.70,label,18,bold=True,align=PP_ALIGN.CENTER)
        if i<5:arrow(s,x+1.75,2.77,x+2.00,2.77)
    text(s,.72,3.90,1.4,.65,'250',34,bold=True,color=ACCENT)
    text(s,2.26,3.93,5.57,.67,'students in the ACCESS\nrandomized trial',20)
    text(s,.72,4.82,7.60,.75,'CBT + mentoring improved\nexecutive functioning.',22,bold=True)
    text(s,.72,5.76,7.52,.60,'Coaching use was also linked to GPA gains\nat a specialist campus.',18,color=MUTED)
    box(s,9.08,3.78,3.58,1.27,line=LINE)
    text(s,9.29,3.96,3.1,.23,'STRENGTH',11,bold=True,color=ACCENT)
    text(s,9.29,4.30,3.13,.61,'Targets planning\nand self-regulation.',18)
    box(s,9.08,5.23,3.58,1.27,line=LINE)
    text(s,9.29,5.41,3.1,.23,'LIMITATION',11,bold=True,color=ACCENT)
    text(s,9.29,5.75,3.13,.60,'Needs trained staff\nand recurring time.',18)
    footer(s,3,['anastopoulos','dupaul'])
    add_notes(s,SCRIPTS[3],'~80 seconds',
        'The diagram is a proposed coaching routine, not the exact ACCESS protocol. ACCESS tested a bundled CBT-plus-mentoring program, not coaching alone. '
        'Its reported d=.39–1.21 spans several outcomes and is intentionally omitted from the slide to avoid implying a GPA or coaching-only effect. '
        'DuPaul et al. used observational service-use data; no precise coaching-hour coefficient is quoted because the full coefficient table was not verified. '
        'https://doi.org/10.1037/ccp0000553 ; https://doi.org/10.1111/ldrp.12143')

    # Content 4: institutional pathway, tradeoffs and open comparison.
    s=prs.slides.add_slide(blank)
    header(s,'04 / Possible solution 2','An integrated ADHD\nsupport system')
    labels=[('Easy access','One clear entry point'),('Accommodations','Individual adjustments'),
            ('Coaching /\nmentoring','Continuing skills support'),('Ongoing\nfollow-up','Review changing needs')]
    for i,(title,detail) in enumerate(labels):
        x=.65+i*3.15
        box(s,x,2.21,2.57,1.39,line=ACCENT)
        text(s,x+.16,2.40,2.27,.66,title,21,bold=True)
        text(s,x+.16,3.22,2.27,.30,detail,12.5,color=MUTED)
        if i<3:arrow(s,x+2.67,2.91,x+3.03,2.91)
    text(s,.73,3.83,11.80,.33,'Faculty awareness  ·  Advising  ·  Mental-health referral when needed',18,color=MUTED)
    text(s,.73,4.41,2.05,.63,'1,980',32,bold=True,color=ACCENT)
    text(s,2.79,4.49,5.27,.53,'students with disabilities',19)
    text(s,.73,5.17,7.60,.72,'Earlier, sustained accommodation use\nwas associated with better outcomes.',19,bold=True)
    text(s,.73,6.08,7.59,.43,'Broad disability sample; observational evidence.',16,color=MUTED)
    box(s,9.08,4.34,3.58,.95,line=LINE)
    text(s,9.28,4.46,3.1,.20,'STRENGTH',10.5,bold=True,color=ACCENT)
    text(s,9.28,4.77,3.15,.41,'Access + continuity',18)
    box(s,9.08,5.47,3.58,.95,line=LINE)
    text(s,9.28,5.59,3.1,.20,'LIMITATION',10.5,bold=True,color=ACCENT)
    text(s,9.28,5.90,3.15,.41,'Coordination across services',17.5)
    footer(s,4,['alvarez','blasey','romhild'])
    add_notes(s,SCRIPTS[4],'~85 seconds, including closing',
        'The service pathway is a proposed design, not a tested Kazakhstan intervention. '
        'Blasey et al.: N=1,980; broad disability sample, not ADHD-only; service use and outcomes are observational. '
        'Römhild & Hollederer review heterogeneous services and outcomes; generic accommodations are not a guaranteed GPA intervention. '
        'No final recommendation is made. End with the stated next research step. '
        'https://doi.org/10.1177/00332941221078011 ; https://doi.org/10.1080/08856257.2023.2195074 ; https://doi.org/10.3389/fpsyg.2023.1216692')

    # Final support slide: all eight sources actually used, alphabetized column-wise.
    s=prs.slides.add_slide(blank)
    text(s,.65,.43,12,.30,'SOURCE LIST / APA 7',11,bold=True,color=ACCENT)
    text(s,.65,.88,12,.63,'References',32,bold=True)
    for key,y in zip(['alvarez','anastopoulos','blasey','dupaul'],[1.83,2.93,4.21,5.33]):
        reference(s,key,.68,y,5.89,1.23 if key in ['anastopoulos','dupaul'] else 1.1)
    for key,y in zip(['gormley','nu','romhild','ross'],[1.83,3.14,4.14,5.57]):
        reference(s,key,6.95,y,5.69,1.4 if key in ['romhild','ross'] else 1.26)
    add_notes(s,'References are provided for questions and follow-up; do not read this slide aloud.',
              'Not part of the timed spoken presentation')
    path=ROOT/'ADHD_Support_Kazakhstan.pptx';prs.save(path)
    return path


def write_support_files(names,course,date):
    words=[len(re.findall(r"\b[\w’–-]+\b",s)) for s in SCRIPTS]
    notes=['# Speaker notes\n','Target: approximately 5–6 minutes, including a brief introduction and closing.\n']
    for i,s in enumerate(SCRIPTS):
        notes.append(f'## Slide {i+1} — {words[i]} spoken words\n\n{s.strip()}\n')
    (ROOT/'SPEAKER_NOTES.md').write_text('\n'.join(notes))
    (ROOT/'build_metadata.json').write_text(json.dumps({'slides':6,'content_slides':4,'font':FONT,
        'presenter':names,'course':course,'date':date,'spoken_words':sum(words),'words_per_slide':words,
        'estimated_minutes_at_130_wpm':round(sum(words)/130,2),'sources':SOURCES},indent=2,ensure_ascii=False))
    (ROOT/'README.md').write_text(f'''# ADHD support in Kazakhstan — topic proposal

Open **ADHD_Support_Kazakhstan.pptx**. It contains six slides: cover, exactly four content slides, and one APA 7 References slide. Text, diagrams and the GPA chart are native, editable PowerPoint objects; the chart includes an embedded Excel data sheet. Speaker notes are embedded and also available in `SPEAKER_NOTES.md`.

The user's build brief takes precedence over the supplied ITMO PDF. The design borrows its large left-aligned titles, whitespace, outlined cards and one accent colour. No artwork or chart screenshot from the PDF is reproduced.

## Graph provenance

`charts/first_year_gpa.png`, `.svg` and `.html` are original Plotly renderings of **Gormley et al. (2019), Table 2**: fall ADHD 2.91, comparison 3.26; spring ADHD 2.79, comparison 3.13. The PowerPoint contains a native editable version of the same grouped bar chart, using the same data in `charts/gormley_2019_table2.csv`. The y-axis is 0–4. Published d values −0.48 and −0.41 are stored for audit but not displayed. No confidence intervals were invented. These descriptive group means are not support-versus-no-support data.

Source DOI: https://doi.org/10.1177/1087054715623046
Primary table: https://pmc.ncbi.nlm.nih.gov/articles/PMC6209537/
Author manuscript: https://libres.uncg.edu/ir/uncg/f/A_Anastopoulos_First_Year_2019.pdf

The mechanism and service/coaching pathways are original diagrams illustrating the proposal, not extracted data. International studies support investigating these approaches; they do not establish their effectiveness in Kazakhstan. The Kazakhstan gap is explicitly a working proposition. Nazarbayev University's documented service is acknowledged. The optional South African secondary finding is omitted to protect the four-slide limit and visual focus; no support variable is invented.

## Timing and cover

Spoken script: {sum(words)} words, approximately {sum(words)/130:.1f} minutes at 130 words/minute, plus short transitions/pauses. Presenter notes distinguish the ACCESS bundled intervention from coaching alone and the broad-disability sample from an ADHD-only sample. There is no final recommendation.

Cover: `{names}` · `{course}` · `{date}`. Bracketed fields are editable placeholders awaiting personal details.

## Fonts and editing

Inter Regular, SemiBold and Bold are packaged in `assets/fonts/` with the SIL Open Font License. Install them on another computer before editing to preserve wrapping. The PDF preview preserves the rendered appearance. The chart, text and shapes remain editable in PowerPoint.

## Rebuild

From the analysis project: `.venv/bin/python presentation/build_presentation.py`

Optional flags: `--presenter 'Name(s)' --course 'Course' --date 'Date'`. The build generates the PowerPoint, chart assets, CSV, speaker-note Markdown and metadata. Requirements are in `requirements.txt`. Plotly PNG/SVG export requires Kaleido and Chrome. Existing project-local Chrome is detected automatically. `render_presentation.applescript` exports the PPTX to a PDF with Keynote for visual checking, without rewriting the original PPTX.

See `SOURCE_AUDIT.md` for the claim-to-source checks and `QA_REPORT.md` for technical and visual validation. The `research/` working cache is excluded from version control; downloaded research papers are not redistributed as presentation assets.
''')


if __name__=='__main__':
    p=argparse.ArgumentParser();p.add_argument('--presenter',default='[Presenter name(s)]')
    p.add_argument('--course',default='[Course]');p.add_argument('--date',default='[Presentation date]')
    a=p.parse_args();chart_assets();path=build(a.presenter,a.course,a.date)
    write_support_files(a.presenter,a.course,a.date);print(path)
